from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(r"C:\Users\PC\Desktop\探矿气象项目代码")
OUT_PATH = PROJECT_ROOT / "docs" / "发明专利技术交底书_气候解耦矿产远景预测.docx"


TITLE = "发明专利技术交底书"
SUBTITLE = "一种面向矿产远景预测的气候影响解耦与目标排序方法、系统及存储介质"


def set_run_font(run, name="宋体", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_cm):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.167

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("发明专利技术交底书 - 气候解耦矿产远景预测")
    set_run_font(run, "宋体", 9, color="666666")


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(TITLE)
    set_run_font(run, "微软雅黑", 22, bold=True, color="0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(SUBTITLE)
    set_run_font(run, "微软雅黑", 13, bold=True, color="1F4D78")

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [4.2, 12.3])
    info = [
        ("技术领域", "矿产资源勘查、地学大数据分析、机器学习、因果推理"),
        ("适用对象", "斑岩铜矿及可扩展至铀矿、铅锌矿、金矿、稀土矿等其他矿种"),
        ("核心方法", "多源空间对齐、气候敏感图、气候影响残差化、Top-K 探矿目标排序"),
        ("主要输出", "气候解耦后的矿产远景预测模型与候选探矿目标优先级列表"),
        ("文档用途", "供发明专利撰写、专利代理沟通和技术方案内部评审使用"),
    ]
    for row, (k, v) in zip(table.rows, info):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        set_cell_shading(row.cells[0], "F2F4F7")
        row.cells[0].paragraphs[0].add_run(k)
        row.cells[1].paragraphs[0].add_run(v)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 10.5, bold=(cell == row.cells[0]))

    doc.add_page_break()


def add_h1(doc, text):
    doc.add_heading(text, level=1)


def add_h2(doc, text):
    doc.add_heading(text, level=2)


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22)
    p.add_run(text)
    for run in p.runs:
        set_run_font(run, "宋体", 11)
    return p


def add_list(doc, items, numbered=True):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)
        for run in p.runs:
            set_run_font(run, "宋体", 11)


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths_cm)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(hdr[idx], "F2F4F7")
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr[idx])
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, "微软雅黑", 10.5, bold=True, color="0B2545")
    for row_data in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run_font(r, "宋体", 10)
    return table


def build_document():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_doc(doc)
    add_title(doc)

    add_h1(doc, "一、发明名称")
    add_para(doc, "一种面向矿产远景预测的气候影响解耦与目标排序方法、系统及存储介质。")

    add_h1(doc, "二、技术领域")
    add_para(doc, "本发明属于矿产资源勘查、地学大数据分析、机器学习和因果推理技术领域，具体涉及一种面向矿产远景预测的多源地学数据融合、气候影响识别、气候敏感特征解耦以及探矿目标排序方法。")

    add_h1(doc, "三、背景技术")
    add_para(doc, "矿产远景预测通常通过融合地质、地球化学、地球物理、地形、遥感等多源数据，识别具有较高成矿潜力的区域。该类方法能够缩小野外勘查范围，降低钻探和实地验证成本。")
    add_para(doc, "现有方法多将不同来源的地学特征直接输入机器学习模型，例如随机森林、梯度提升树、逻辑回归或深度学习模型，用于预测某一区域是否具有矿化潜力。然而，在实际探矿场景中，地表观测特征并不只受成矿作用影响，还可能受到气候、风化、侵蚀、植被覆盖、地表暴露程度、人类活动等因素影响。")
    add_para(doc, "例如，地球化学异常可能受到降水、温度、蒸散、土壤湿度和风化强度的影响；遥感光谱特征可能受到植被、水分和地表覆盖条件的影响；地形和地表暴露条件也会改变矿化信息在地表的可观测程度。因此，模型可能学习到由气候条件造成的表观差异，而不是稳定的成矿相关信号。")
    add_para(doc, "现有矿产远景预测方法通常存在以下不足：")
    add_list(doc, [
        "多源地学特征直接融合，缺少对气候影响和地表观测偏差的建模。",
        "缺少显式的大气或气候变量参与探矿特征分析。",
        "缺少识别气候敏感探矿特征的方法。",
        "缺少在模型训练前对气候影响进行解耦或残差化处理的方法。",
        "多数方法侧重整体分类指标，而实际探矿更关注有限预算下排名靠前的目标点质量。",
    ])

    add_h1(doc, "四、发明目的")
    add_para(doc, "本发明的目的在于提供一种面向矿产远景预测的气候影响解耦与目标排序方法，用于解决现有方法中多源特征受气候因素混杂、模型可能学习到非成矿相关表观模式、以及预测结果难以服务于有限预算探矿决策的问题。")
    add_para(doc, "本发明希望实现以下目标：")
    add_list(doc, [
        "构建包含矿点标签、地球化学、地球物理、地质、断层、地形、气候等信息的多源对齐样本数据集。",
        "识别受气候影响较强的探矿特征。",
        "构建气候敏感图，用于刻画气候变量与探矿特征之间的关联关系。",
        "对气候敏感探矿特征进行气候影响解耦。",
        "使用解耦后的特征训练矿产远景预测模型。",
        "输出候选探矿目标的排序结果，并通过 Top-K 指标评价模型在有限勘查预算下的实际价值。",
    ])

    add_h1(doc, "五、技术方案")
    add_h2(doc, "步骤一：构建多源地学样本数据集")
    add_para(doc, "获取研究区内的矿点数据和非矿点候选样本数据。矿点数据包括已知矿床、矿点或矿化异常点的经纬度、矿种、矿床类型、矿化模型等信息。非矿点候选样本可包括已知矿区邻近负样本、中性样本或研究区背景样本。")
    add_para(doc, "对每个样本点，根据经纬度进行空间对齐，提取以下多源特征：")
    add_table(doc, ["数据类别", "主要内容"], [
        ("地球化学特征", "铜、钼、银、砷、锑、铋、铅、锌等元素浓度或异常指标。"),
        ("地球物理特征", "重力异常、重力导数、磁异常或相关派生特征。"),
        ("地质特征", "岩性、地层、侵入岩、构造背景等。"),
        ("断层与构造特征", "断层距离、断层密度、断层类型等。"),
        ("地形特征", "高程、坡度、地形起伏度等。"),
        ("气候特征", "温度、降水、蒸散、太阳辐射、土壤水分、雪水当量、气候水分亏缺等。"),
        ("可选遥感特征", "多光谱反射率、植被指数、短波红外指标等。"),
    ], [4.2, 12.3])
    add_para(doc, "对上述特征进行缺失值处理、异常值截断、稳健标准化和空间聚合，形成统一的样本特征表。")

    add_h2(doc, "步骤二：划分样本标签")
    add_para(doc, "根据研究目标构建正样本、负样本和中性样本。正样本为目标矿种或目标矿床类型的已知矿点。负样本为已知矿区附近一定距离范围内、但不属于目标矿种或目标矿床类型的地点，或根据研究设定从已知非目标矿区中抽取的地点。中性样本为研究区内尚未探测到目标矿点的背景地点。中性样本可用于辅助模型理解背景分布，也可根据任务设定不参与监督训练，仅参与候选排序或推理。")

    add_h2(doc, "步骤三：构建气候敏感图")
    add_para(doc, "针对多源特征中的气候变量和探矿特征，构建气候敏感图。气候敏感图中的节点包括气候概念节点和探矿特征节点。气候概念节点可包括温度、降水、蒸散、干旱程度、雪影响、辐射、水分条件等。探矿特征节点可包括地球化学异常、地球物理异常、地形暴露特征、断层构造特征和遥感地表特征等。")
    add_para(doc, "气候敏感图的边用于表示气候变量与探矿特征之间可能存在的影响关系。该关系可由以下证据综合获得：")
    add_list(doc, [
        "全局气候-特征统计关联，例如 Spearman 相关系数。",
        "交叉验证训练折内的局部气候敏感性。",
        "不同环境分区下稳定出现的候选因果边。",
        "领域知识约束，例如气候变量可影响地表观测特征，但矿点标签不作为气候变量的原因。",
    ])
    add_para(doc, "根据上述证据计算综合气候敏感得分。当某一探矿特征与气候变量之间的综合得分超过预设阈值时，将该探矿特征标记为气候敏感特征。")

    add_h2(doc, "步骤四：对气候敏感特征进行解耦处理")
    add_para(doc, "对识别出的气候敏感探矿特征进行气候影响解耦。具体方法为：以气候变量作为解释变量，以某一气候敏感探矿特征作为被解释变量，在训练集内拟合气候影响模型。然后使用该模型估计气候变量对该探矿特征的可解释部分，并从原始特征中去除该部分，得到残差化后的探矿特征。")
    add_para(doc, "残差化后的特征可表示为：X_residual = X_original - X_climate_hat。其中，X_original 为原始探矿特征，X_climate_hat 为由气候变量预测得到的气候相关部分，X_residual 为去除气候影响后的特征。该步骤可减少模型对气候条件的依赖，使模型更关注稳定的成矿相关信号。")

    add_h2(doc, "步骤五：训练矿产远景预测模型")
    add_para(doc, "使用解耦后的多源特征训练矿产远景预测模型。可选模型包括逻辑回归、随机森林、梯度提升树、直方图梯度提升树、XGBoost、LightGBM 或其他机器学习模型。")
    add_para(doc, "训练过程中，可比较以下特征组合：")
    add_list(doc, [
        "全特征模型：直接使用所有特征，包括气候变量。",
        "无气候模型：去除气候变量，仅使用非气候特征。",
        "气候解耦模型：对气候敏感特征进行残差化后训练。",
        "图引导气候解耦模型：根据气候敏感图选择特征并进行残差化后训练。",
    ])

    add_h2(doc, "步骤六：输出探矿目标排序结果")
    add_para(doc, "模型输出每个候选样本点属于目标矿种或目标矿床类型的概率分数。根据概率分数对候选样本进行降序排序，得到探矿优先级列表。由于实际探矿通常受钻探预算和实地验证成本约束，本发明采用 Top-K 推荐指标评价模型效果，包括 Precision@K、Recall@K、F1@K、Lift@K、NDCG@K、PR-AUC 和 ROC-AUC。Top-K 指标用于衡量模型是否能将真实矿点排在候选列表前部，更符合实际探矿决策需求。")

    add_h1(doc, "六、有益效果")
    add_para(doc, "与现有技术相比，本发明具有以下有益效果：")
    add_list(doc, [
        "将气候变量显式引入矿产远景预测流程，能够分析气候对地表探矿特征的影响。",
        "构建气候敏感图，用于识别受气候影响较强的探矿特征。",
        "通过残差化方法去除探矿特征中的气候相关部分，降低气候混杂因素对模型判断的影响。",
        "能够提升候选矿点排序质量，尤其适用于钻探预算有限、只关注前若干高优先级目标的实际勘查场景。",
        "可扩展至不同矿种、不同区域和不同数据组合，具有较强的通用性。",
        "不仅适用于斑岩铜矿，也可扩展至铀矿、铅锌矿、金矿、稀土矿等其他矿产资源远景预测任务。",
    ])

    add_h1(doc, "七、具体实施例")
    add_para(doc, "以美国西部地区斑岩铜矿为例，构建正样本、负样本和中性样本。正样本来自已知斑岩铜矿点。负样本来自已知矿区附近但不属于斑岩铜矿的地点。中性样本来自研究区内尚未探测到斑岩铜矿的背景区域。")
    add_para(doc, "对每个样本点提取地球化学、地球物理、地质、断层、地形和气候特征。气候特征可来自 TerraClimate 或其他气候数据源，地形特征可来自 DEM，地球物理特征可来自重力异常或磁异常数据，地质和断层特征可来自 USGS 地质图和断层数据。")
    add_para(doc, "随后，计算气候变量与探矿特征之间的统计关联，并结合不同环境分区下稳定出现的候选因果关系，构建气候敏感图。对气候敏感图中连接较强的探矿特征进行残差化处理，得到气候解耦后的特征表。")
    add_para(doc, "最后，使用随机森林、梯度提升树等模型进行训练，并对候选点进行排序。实验结果表明，经过气候解耦处理后，模型在 Top-5% 或 Top-10% 的候选目标排序中能够获得更高的 Precision、Recall、F1、Lift 或 NDCG 指标，说明该方法能够提高有限预算下的高价值探矿目标识别能力。")

    add_h1(doc, "八、可替代方案")
    add_para(doc, "本发明中的气候敏感图可以采用不同方式构建，包括但不限于：")
    add_list(doc, [
        "基于 Spearman 相关系数构建。",
        "基于偏相关分析构建。",
        "基于 PC 算法、GES 算法、NOTEARS 等因果发现算法构建。",
        "基于专家知识规则构建。",
        "基于机器学习模型特征重要性构建。",
        "基于多种证据融合构建。",
    ])
    add_para(doc, "本发明中的气候解耦方法也可以采用不同形式，包括但不限于线性回归残差化、Ridge 或 Lasso 残差化、随机森林回归残差化、梯度提升回归残差化、神经网络残差化、反事实样本增强和域不变表示学习。本发明中的目标排序模型也不限于随机森林，可替换为任意分类模型、排序模型或深度学习模型。")

    add_h1(doc, "九、拟保护的核心创新点")
    add_list(doc, [
        "一种将气候变量引入矿产远景预测并识别气候敏感探矿特征的方法。",
        "一种基于气候敏感图指导探矿特征解耦的方法。",
        "一种对地球化学、地球物理、地质、断层、地形、气候等多源特征进行统一空间对齐并用于矿产预测的方法。",
        "一种通过残差化去除探矿特征中气候相关部分的方法。",
        "一种面向有限钻探预算的矿产远景目标排序方法。",
        "一种适用于不同矿种和不同区域的气候解耦矿产远景预测系统。",
    ])

    add_h1(doc, "十、建议权利要求方向")
    add_h2(doc, "独立权利要求一：方法权利要求")
    add_para(doc, "一种面向矿产远景预测的气候影响解耦与目标排序方法，其特征在于，包括：获取研究区内矿点样本和候选样本；根据样本经纬度对齐多源地学数据，形成样本特征表；从所述样本特征表中识别气候变量和探矿特征；根据气候变量与探矿特征之间的关联关系构建气候敏感图；根据所述气候敏感图确定气候敏感探矿特征；利用气候变量对所述气候敏感探矿特征进行残差化处理，得到气候解耦特征；基于所述气候解耦特征训练矿产远景预测模型；利用所述矿产远景预测模型输出候选样本的矿产远景分数，并根据所述矿产远景分数生成探矿目标排序结果。")

    add_h2(doc, "独立权利要求二：系统权利要求")
    add_para(doc, "一种面向矿产远景预测的气候影响解耦与目标排序系统，包括数据获取模块、空间对齐模块、气候敏感图构建模块、气候敏感特征识别模块、气候影响解耦模块、矿产远景预测模块和目标排序模块。其中，所述气候影响解耦模块用于根据气候敏感图对气候敏感探矿特征进行残差化处理。")

    add_h2(doc, "独立权利要求三：存储介质权利要求")
    add_para(doc, "一种计算机可读存储介质，其上存储有计算机程序，所述程序被处理器执行时实现上述面向矿产远景预测的气候影响解耦与目标排序方法。")

    add_h1(doc, "十一、附图建议")
    add_list(doc, [
        "方法总体流程图。",
        "多源数据空间对齐示意图。",
        "气候敏感图构建示意图。",
        "气候敏感特征残差化处理示意图。",
        "探矿目标排序输出示意图。",
        "系统模块结构图。",
    ])

    add_h1(doc, "十二、摘要建议")
    add_para(doc, "本发明公开了一种面向矿产远景预测的气候影响解耦与目标排序方法、系统及存储介质。该方法首先获取矿点样本和候选样本，并根据经纬度对齐地球化学、地球物理、地质、断层、地形和气候等多源数据，构建统一样本特征表；然后根据气候变量与探矿特征之间的关联关系构建气候敏感图，并利用该气候敏感图识别受气候影响的探矿特征；随后利用气候变量对气候敏感探矿特征进行残差化处理，得到气候解耦后的特征；最后基于解耦特征训练矿产远景预测模型，并输出候选探矿目标排序结果。本发明能够降低气候因素对地表探矿特征的混杂影响，提高有限勘查预算下高优先级探矿目标的识别能力，适用于斑岩铜矿及其他矿种的矿产远景预测任务。")

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
