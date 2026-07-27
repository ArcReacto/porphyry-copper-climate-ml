from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE_MD = ROOT / "docs" / "今日优化实验步骤与结果汇总.md"
OUT_DOCX = ROOT / "docs" / "模型说明与实验结果表.docx"


INCLUDE_HEADERS = {
    ("类别", "模型", "模型定位", "说明"),
    ("项目", "数量"),
    ("目标概念", "类型", "主要关联气候概念", "后续用途"),
    ("模型", "模型定位", "模型角色", "具体含义"),
    ("模型", "模型定位", "模型角色", "ROC-AUC", "AP/AUPRC", "Balanced Accuracy", "F1", "Top-5% Recall", "Top-10% Recall"),
    ("模型", "模型定位", "模型角色", "mean abs delta", "p95 abs delta", "frac abs delta > 0.05", "AP 变化", "F1 变化"),
    ("模型", "模型定位", "模型角色", "训练方式", "ROC-AUC", "AP/AUPRC", "Balanced Accuracy", "F1", "Top-5% Recall", "Top-10% Recall"),
    ("模型", "模型定位", "模型角色", "ROC-AUC", "AP/AUPRC", "Balanced Accuracy", "Precision", "Recall", "F1", "Top-5% Recall", "Top-10% Recall"),
    ("指标", "最佳模型", "模型定位", "模型角色", "数值"),
    ("模型", "是否基线", "论文角色"),
}

SKIP_IF_HEADERS = {
    ("脚本", "功能"),
    ("指标", "作用"),
    ("场景", "含义"),
}


def clean_cell(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^\*\*(.*)\*\*$", r"\1", text)
    text = text.replace("`", "")
    text = text.replace("<br>", "\n")
    return text


def parse_markdown_tables(md: str) -> list[dict]:
    lines = md.splitlines()
    current_headings: list[str] = []
    last_nonempty = ""
    tables: list[dict] = []
    i = 0

    while i < len(lines):
        line = lines[i]
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            current_headings = current_headings[: level - 1] + [title]
            last_nonempty = title
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[i + 1]):
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1

            rows = []
            for row_line in table_lines:
                cells = [clean_cell(cell) for cell in row_line.strip().strip("|").split("|")]
                rows.append(cells)

            header = tuple(rows[0])
            if header not in SKIP_IF_HEADERS and header in INCLUDE_HEADERS:
                title = " / ".join(current_headings)
                if last_nonempty and not last_nonempty.startswith("#") and last_nonempty != current_headings[-1]:
                    title = f"{title} - {last_nonempty.rstrip('：:')}"
                if header == ("模型", "是否基线", "论文角色"):
                    title = "当前论文主线建议"
                tables.append(
                    {
                        "title": title,
                        "header": rows[0],
                        "rows": rows[2:],
                    }
                )
            continue

        if line.strip() and line.strip() != "```":
            last_nonempty = line.strip()
        i += 1
    return tables


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False, font_size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_row_cannot_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_header_row_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, title: str, header: list[str], rows: list[list[str]]) -> None:
    page_break_before_tokens = [
        "主要选中的目标概念",
        "6.1 模型说明",
        "8.1 模型说明",
        "9.1 主性能表",
    ]
    if any(token in title for token in page_break_before_tokens):
        doc.add_page_break()

    heading = doc.add_heading(title, level=2)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.keep_together = True
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    for cell, text in zip(table.rows[0].cells, header):
        set_cell_shading(cell, "E8EEF5")
        set_cell_text(cell, text, bold=True, font_size=7 if len(header) >= 9 else 8)
    set_header_row_repeat(table.rows[0])
    set_row_cannot_split(table.rows[0])

    for row in rows:
        cells = table.add_row().cells
        for idx, cell in enumerate(cells):
            value = row[idx] if idx < len(row) else ""
            set_cell_text(cell, value, font_size=7 if len(header) >= 9 else 8)
        set_row_cannot_split(table.rows[-1])

    doc.add_paragraph()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10)
    for style_name, size, color in [
        ("Title", 20, "1F4D78"),
        ("Heading 1", 15, "2E74B5"),
        ("Heading 2", 12, "2E74B5"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def main() -> None:
    md = SOURCE_MD.read_text(encoding="utf-8")
    tables = parse_markdown_tables(md)
    if not tables:
        raise RuntimeError(f"No selected tables were parsed from {SOURCE_MD}")

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("模型说明与实验结果表")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("斑岩铜探矿气候解耦优化实验汇总").italic = True

    note = doc.add_paragraph()
    note.add_run("说明：").bold = True
    note.add_run(
        "本文档从《今日优化实验步骤与结果汇总.md》中单独抽取模型定位、模型说明、"
        "GroupKFold 结果、气候扰动稳定性、反事实增强和最终主实验结果表，便于后续论文写作和汇报复用。"
    )

    doc.add_heading("模型说明与结果表", level=1)
    for table in tables:
        add_table(doc, table["title"], table["header"], table["rows"])

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
