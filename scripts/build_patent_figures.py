from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = PROJECT_ROOT / "docs" / "patent_figures"
DOCX_PATH = PROJECT_ROOT / "docs" / "专利附图_气候解耦矿产远景预测.docx"

FONT_REG = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"

W, H = 1800, 1100
INK = (24, 24, 24)
MUTED = (105, 105, 105)
LINE = (35, 35, 35)
FILL = (248, 248, 248)
FILL2 = (235, 239, 244)
WHITE = (255, 255, 255)


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size=size)


def set_docx_run_font(run, east_asia="微软雅黑", size=None, bold=None):
    run.font.name = east_asia
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


F_TITLE = font(42, True)
F_H = font(30, True)
F_BODY = font(25)
F_SMALL = font(21)
F_TINY = font(18)


def new_canvas(title):
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    d.text((W // 2, 45), title, font=F_TITLE, fill=INK, anchor="mm")
    d.line((110, 96, W - 110, 96), fill=(160, 160, 160), width=2)
    return img, d


def text_size(d, text, fnt):
    box = d.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_zh(text, max_chars):
    lines = []
    for seg in str(text).split("\n"):
        if len(seg) <= max_chars:
            lines.append(seg)
        else:
            lines.extend(textwrap.wrap(seg, width=max_chars, break_long_words=True, replace_whitespace=False))
    return lines


def draw_center_text(d, box, text, fnt=F_BODY, fill=INK, max_chars=12, line_gap=8):
    x1, y1, x2, y2 = box
    lines = wrap_zh(text, max_chars)
    heights = [text_size(d, line, fnt)[1] for line in lines]
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, h in zip(lines, heights):
        d.text(((x1 + x2) / 2, y), line, font=fnt, fill=fill, anchor="ma")
        y += h + line_gap


def draw_box(d, box, title, body=None, fill=FILL, outline=LINE, width=3, max_chars=12):
    d.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = box
    if body is None:
        draw_center_text(d, box, title, F_BODY, max_chars=max_chars)
    else:
        d.text(((x1 + x2) / 2, y1 + 25), title, font=F_H, fill=INK, anchor="ma")
        lines = wrap_zh(body, max_chars)
        y = y1 + 70
        for line in lines:
            d.text(((x1 + x2) / 2, y), line, font=F_SMALL, fill=MUTED, anchor="ma")
            y += 32


def arrow(d, start, end, width=4, fill=LINE):
    x1, y1 = start
    x2, y2 = end
    d.line((x1, y1, x2, y2), fill=fill, width=width)
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 18
    pts = [
        (x2, y2),
        (x2 - size * math.cos(ang - math.pi / 6), y2 - size * math.sin(ang - math.pi / 6)),
        (x2 - size * math.cos(ang + math.pi / 6), y2 - size * math.sin(ang + math.pi / 6)),
    ]
    d.polygon(pts, fill=fill)


def dashed_rect(d, box, label):
    x1, y1, x2, y2 = box
    dash = 16
    gap = 10
    for x in range(x1, x2, dash + gap):
        d.line((x, y1, min(x + dash, x2), y1), fill=(110, 110, 110), width=2)
        d.line((x, y2, min(x + dash, x2), y2), fill=(110, 110, 110), width=2)
    for y in range(y1, y2, dash + gap):
        d.line((x1, y, x1, min(y + dash, y2)), fill=(110, 110, 110), width=2)
        d.line((x2, y, x2, min(y + dash, y2)), fill=(110, 110, 110), width=2)
    d.text((x1 + 22, y1 - 30), label, font=F_SMALL, fill=MUTED)


def save(img, name):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    img.save(path, dpi=(300, 300))
    return path


def fig1_method_flow():
    img, d = new_canvas("图1  方法总体流程图")
    steps = [
        ("S101", "获取样本", "矿点、负样本、中性样本"),
        ("S102", "空间对齐", "多源地学与气候数据"),
        ("S103", "标签构建", "正负样本与候选样本"),
        ("S104", "气候敏感图", "识别气候敏感特征"),
        ("S105", "气候解耦", "去除气候相关部分"),
        ("S106", "模型训练", "矿产远景预测模型"),
        ("S107", "目标排序", "输出 Top-K 探矿目标"),
    ]
    x0, y0 = 110, 190
    bw, bh, gap = 210, 170, 32
    boxes = []
    for i, (sid, title, body) in enumerate(steps):
        row = 0 if i < 4 else 1
        col = i if i < 4 else 6 - i
        x = x0 + col * (bw + gap)
        y = y0 + row * 330
        boxes.append((x, y, x + bw, y + bh))
        draw_box(d, boxes[-1], f"{sid}\n{title}", body, max_chars=10)
    for i in range(3):
        arrow(d, (boxes[i][2] + 5, (boxes[i][1] + boxes[i][3]) // 2), (boxes[i + 1][0] - 8, (boxes[i + 1][1] + boxes[i + 1][3]) // 2))
    arrow(d, ((boxes[3][0] + boxes[3][2]) // 2, boxes[3][3] + 6), ((boxes[4][0] + boxes[4][2]) // 2, boxes[4][1] - 8))
    for i in range(4, 6):
        arrow(d, (boxes[i][0] - 5, (boxes[i][1] + boxes[i][3]) // 2), (boxes[i + 1][2] + 8, (boxes[i + 1][1] + boxes[i + 1][3]) // 2))
    dashed_rect(d, (75, 150, 1725, 845), "基于多源数据的气候影响解耦与目标排序流程")
    return save(img, "fig1_method_flow.png")


def fig2_spatial_alignment():
    img, d = new_canvas("图2  多源数据空间对齐示意图")
    left_items = ["矿点坐标", "地球化学", "地球物理", "地质与断层", "DEM 地形", "气候数据"]
    y = 175
    for item in left_items:
        draw_box(d, (105, y, 390, y + 90), item, fill=WHITE, max_chars=9)
        y += 115
    draw_box(d, (610, 230, 1025, 690), "经纬度空间索引", "缓冲区 / 最近邻 / 栅格采样 / 邻域统计", fill=FILL2, max_chars=12)
    for yy in [220, 335, 450, 565, 680, 795]:
        arrow(d, (390, yy), (600, 460))
    draw_box(d, (1210, 230, 1660, 690), "统一样本特征表", "每一行 = 一个样本点\n每一列 = 一个对齐后的特征", fill=WHITE, max_chars=16)
    arrow(d, (1030, 460), (1200, 460))
    d.ellipse((725, 420, 910, 605), outline=LINE, width=3)
    d.ellipse((760, 455, 875, 570), outline=LINE, width=2)
    d.ellipse((805, 500, 830, 525), fill=INK)
    d.text((817, 548), "样本点", font=F_TINY, fill=MUTED, anchor="mm")
    return save(img, "fig2_spatial_alignment.png")


def fig3_climate_sensitivity_graph():
    img, d = new_canvas("图3  气候敏感图构建示意图")
    evidences = [
        "全局气候-特征关联",
        "训练折内局部敏感性",
        "跨环境稳定候选因果边",
        "领域知识方向约束",
    ]
    y = 195
    for e in evidences:
        draw_box(d, (100, y, 430, y + 95), e, fill=WHITE, max_chars=12)
        arrow(d, (430, y + 47), (590, 455))
        y += 130
    draw_box(d, (600, 310, 905, 600), "证据融合", "综合气候敏感得分\nScore >= 阈值", fill=FILL2, max_chars=12)
    arrow(d, (910, 455), (1050, 455))
    dashed_rect(d, (1070, 185, 1665, 750), "气候敏感图")
    climates = [("温度", 1160, 285), ("降水", 1160, 430), ("蒸散", 1160, 575)]
    features = [("地球化学异常", 1510, 250), ("重力源强度", 1510, 365), ("地形暴露", 1510, 480), ("光谱蚀变特征", 1510, 595)]
    for t, x, y in climates:
        d.ellipse((x - 70, y - 45, x + 70, y + 45), fill=WHITE, outline=LINE, width=3)
        draw_center_text(d, (x - 70, y - 45, x + 70, y + 45), t, F_SMALL, max_chars=5)
    for t, x, y in features:
        draw_box(d, (x - 105, y - 45, x + 105, y + 45), t, fill=FILL, max_chars=7)
    for sx, sy, ex, ey in [
        (1230, 285, 1405, 250), (1230, 430, 1405, 365), (1230, 430, 1405, 480),
        (1230, 575, 1405, 480), (1230, 285, 1405, 595), (1230, 575, 1405, 595),
    ]:
        arrow(d, (sx, sy), (ex, ey), width=3)
    draw_box(d, (635, 735, 1390, 835), "输出：气候敏感特征清单及对应气候影响边", fill=WHITE, max_chars=25)
    arrow(d, (1340, 750), (1340, 705))
    return save(img, "fig3_climate_sensitivity_graph.png")


def fig4_residualization():
    img, d = new_canvas("图4  气候敏感特征残差化处理示意图")
    draw_box(d, (110, 220, 430, 390), "气候变量 C", "温度、降水、蒸散、水分等", fill=WHITE, max_chars=11)
    draw_box(d, (110, 600, 430, 770), "原始探矿特征 X", "地化异常、重力、光谱等", fill=WHITE, max_chars=12)
    draw_box(d, (650, 250, 1030, 540), "气候影响模型", "在训练集内拟合\nX = f(C) + ε", fill=FILL2, max_chars=13)
    draw_box(d, (650, 610, 1030, 780), "估计气候相关部分", "X_climate_hat = f(C)", fill=WHITE, max_chars=15)
    draw_box(d, (1260, 410, 1650, 610), "解耦后特征", "X_residual = X_original - X_climate_hat", fill=FILL, max_chars=18)
    arrow(d, (430, 305), (645, 360))
    arrow(d, (430, 685), (645, 690))
    arrow(d, (840, 545), (840, 605))
    arrow(d, (1030, 695), (1250, 535))
    arrow(d, (430, 685), (1250, 505))
    d.text((900, 865), "仅使用训练折拟合气候影响模型，避免测试集信息泄漏", font=F_SMALL, fill=MUTED, anchor="mm")
    return save(img, "fig4_residualization.png")


def fig5_target_ranking():
    img, d = new_canvas("图5  探矿目标排序输出示意图")
    draw_box(d, (110, 270, 420, 620), "候选样本点", "已知矿区邻近点\n中性背景点\n待评价目标点", fill=WHITE, max_chars=10)
    draw_box(d, (610, 270, 940, 620), "矿产远景预测模型", "输出每个样本的\n成矿潜力分数", fill=FILL2, max_chars=11)
    draw_box(d, (1130, 185, 1605, 720), "排序结果", "1  目标点 A  0.982\n2  目标点 B  0.971\n3  目标点 C  0.953\n...\nTop-K 高优先级目标", fill=WHITE, max_chars=18)
    arrow(d, (420, 445), (600, 445))
    arrow(d, (940, 445), (1120, 445))
    draw_box(d, (425, 780, 1485, 890), "面向有限钻探预算的评价指标：Precision@K、Recall@K、F1@K、Lift@K、NDCG@K、PR-AUC", fill=FILL, max_chars=35)
    arrow(d, (1365, 720), (1365, 775))
    return save(img, "fig5_target_ranking.png")


def fig6_system_modules():
    img, d = new_canvas("图6  系统模块结构图")
    modules = [
        ("数据获取模块", "矿点、地化、地物、地质、地形、气候"),
        ("空间对齐模块", "经纬度索引与邻域统计"),
        ("气候敏感图构建模块", "融合关联、敏感性与稳定边"),
        ("气候影响解耦模块", "残差化气候敏感特征"),
        ("矿产远景预测模块", "训练分类或排序模型"),
        ("目标排序模块", "生成候选目标优先级"),
    ]
    positions = [(120, 210), (590, 210), (1060, 210), (120, 590), (590, 590), (1060, 590)]
    boxes = []
    for (title, body), (x, y) in zip(modules, positions):
        box = (x, y, x + 360, y + 180)
        boxes.append(box)
        draw_box(d, box, title, body, fill=WHITE, max_chars=13)
    arrow(d, (boxes[0][2], 300), (boxes[1][0], 300))
    arrow(d, (boxes[1][2], 300), (boxes[2][0], 300))
    arrow(d, ((boxes[2][0] + boxes[2][2]) // 2, boxes[2][3]), ((boxes[5][0] + boxes[5][2]) // 2, boxes[5][1] - 10))
    arrow(d, (boxes[3][2], 680), (boxes[4][0], 680))
    arrow(d, (boxes[4][2], 680), (boxes[5][0], 680))
    arrow(d, ((boxes[1][0] + boxes[1][2]) // 2, boxes[1][3]), ((boxes[4][0] + boxes[4][2]) // 2, boxes[4][1] - 10))
    draw_box(d, (635, 850, 1135, 950), "结果存储与评价模块", "保存特征表、模型、Top-K 指标与附图", fill=FILL2, max_chars=18)
    arrow(d, ((boxes[5][0] + boxes[5][2]) // 2, boxes[5][3]), (885, 845))
    dashed_rect(d, (80, 165, 1720, 990), "气候解耦矿产远景预测系统")
    return save(img, "fig6_system_modules.png")


def build_docx(paths):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("专利附图：气候解耦矿产远景预测")
    set_docx_run_font(r, size=18, bold=True)

    captions = [
        "图1  方法总体流程图",
        "图2  多源数据空间对齐示意图",
        "图3  气候敏感图构建示意图",
        "图4  气候敏感特征残差化处理示意图",
        "图5  探矿目标排序输出示意图",
        "图6  系统模块结构图",
    ]
    for i, (path, cap) in enumerate(zip(paths, captions)):
        if i:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(cap)
        set_docx_run_font(r, size=12, bold=True)
        pic_p = doc.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_p.add_run().add_picture(str(path), width=Inches(7.0))
    doc.save(DOCX_PATH)


def main():
    paths = [
        fig1_method_flow(),
        fig2_spatial_alignment(),
        fig3_climate_sensitivity_graph(),
        fig4_residualization(),
        fig5_target_ranking(),
        fig6_system_modules(),
    ]
    build_docx(paths)
    print("Generated:")
    for p in paths:
        print(p)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
